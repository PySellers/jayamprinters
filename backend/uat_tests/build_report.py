"""
Builds the client-facing test report (Excel) from the API and browser test results.

    python build_report.py <api results.json> <ui results glob, e.g. 'ui_results_*.json'> <output.xlsx> [screenshots dir]
"""

import glob
from collections import OrderedDict
import json
import os
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

api_file, ui_file, out = sys.argv[1:4]
shots = sys.argv[4] if len(sys.argv) > 4 else None
api = json.load(open(api_file, encoding="utf-8"))
# Several browser runs may overlap (a re-run of one suite after a fix); the newest run of a scenario wins.
_latest = OrderedDict()
for f in sorted(glob.glob(ui_file), key=os.path.getmtime):
    for r in json.load(open(f, encoding="utf-8")):
        _latest[(r["module"], r["scenario"])] = r
ui = sorted(_latest.values(), key=lambda r: (r["id"].rsplit("-", 1)[0], r["id"]))

HEAD = PatternFill("solid", fgColor="1F3A5F")
GREEN = PatternFill("solid", fgColor="D9F2D9")
RED = PatternFill("solid", fgColor="F8D0D0")
GREY = PatternFill("solid", fgColor="EEF1F5")
WRAP = Alignment(wrap_text=True, vertical="top")

DEFECTS = [
    ("Critical", "Invoices, job cards, purchases and orders", "After any record was deleted, the next invoice / job card / purchase number could repeat an existing one and the save failed with a server error.", "Numbers now continue from the highest number in use, with a lock so two people saving at once cannot clash."),
    ("High", "Sign-in", "A user switched off (deactivated) by the admin could still sign in with their password.", "Deactivated users are refused at sign-in; existing sessions were already blocked."),
    ("High", "Invoices", "A payment larger than the balance due was accepted (e.g. Rs 759.60 recorded against a Rs 259.60 bill).", "Payments above the balance are refused with a clear message."),
    ("High", "Purchases / vendor payments", "Zero or negative quantities and prices were accepted, and a vendor could be paid more than the purchase total.", "Quantity must be above zero, price cannot be negative, and vendor payments cannot exceed the balance."),
    ("High", "Lists (Invoices, Job Cards, Customers, Purchases, Vendors, Users)", "Screens silently showed only the newest 100 records, so older invoices and customers disappeared and dashboard counts stopped at 100.", "Lists now return every record (up to 5,000); invoices stay newest-first."),
    ("High", "Book Covers pricing", "Readymade Cover held Offset, Screen and Digital rates for the same size and quantity with nothing to choose between, so the cheaper Offset rate was always quoted and Screen / Digital could never be reached.", "Split into Readymade Cover - Offset / - Screen / - Digital; each prices only its own quantity range."),
    ("High", "Complete Order screen", "The order screen showed no price at all until the invoice had been generated.", "Added a live rate / amount per line and an order total panel (subtotal, GST, total) that updates as options are chosen."),
    ("High", "Payments and money boxes", "Amounts with paise (e.g. Rs 891.50) could not be entered in the payment box and other money / quantity boxes: the browser silently refused to submit.", "Money and quantity boxes now accept decimals everywhere."),
    ("Low", "Record Payment box", "Entering an amount above the balance did nothing and gave no message.", "The box now shows the balance due and a clear message."),
    ("Medium", "Pricing", "An extra charge that had been switched off could still be billed if its id was sent.", "Inactive charges are refused."),
    ("Medium", "Orders", "An order with no items was accepted, and an order for a customer that does not exist caused a server error.", "Both are refused with a clear message (422 / 404)."),
    ("Medium", "Cash ledger and cheques", "Cash entries and cheques with zero or negative amounts were accepted.", "Amounts must be above zero."),
    ("Medium", "Deleting records", "Deleting a customer who has orders, an invoiced quotation, or a price band / option / extra charge / attribute that is in use caused a server error.", "Clear message explaining why it cannot be deleted and what to do instead."),
    ("Medium", "Products", "Deleting an unused item that had a price grid was blocked with a misleading 'in use' message.", "Unused items delete together with their price grid; items used on orders are protected."),
    ("Low", "Money totals", "Tiny floating-point residue (e.g. 220.00000000000003) could show a fully paid cash bill as part-paid.", "Totals are rounded to the paisa and compared in paise."),
    ("Low", "Price import", "The price sync tool did not copy each item's pricing type, leaving one invitation item on the old type in production.", "Pricing type is now synced too."),
    ("Medium", "Purchases - supplier bill upload (new)", "The first build sent the bill photo / PDF to the server in the wrong format, so every upload was refused with 'Field required'.", "Uploads are now sent as proper file uploads; checked with a real upload in the browser."),
    ("Medium", "Pricing - bands that overlap (new)", "Printout and Scanning have overlapping page bands; the 'a bigger order never costs less' rule could pick the wrong neighbouring band, and which one depended on database order.", "The rule now uses the neighbouring band that actually has a rate for that item; tests replay every price cell."),
    ("Low", "Copies per set (new)", "The new '1' choice was listed last instead of first, because options were shown in the order they were created.", "Options are returned in their display order everywhere."),
    ("Low", "Accessibility", "Edit / Delete icon buttons in the tables had no names for screen readers or automated tests.", "Added Edit / Delete labels."),
]


def style_header(ws, widths, row=1):
    for c in ws[row]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    ws.auto_filter.ref = f"A{row}:{get_column_letter(len(widths))}{ws.max_row}"


def result_sheet(wb, title, rows):
    ws = wb.create_sheet(title)
    ws.append(["Test ID", "Area", "What was tested", "Expected result", "Actual result", "Result"])
    for r in rows:
        ws.append([r["id"], r["module"], r["scenario"], r["expected"], r["actual"], r["status"]])
        ws.cell(ws.max_row, 6).fill = GREEN if r["status"] == "PASS" else RED
        for c in ws[ws.max_row]:
            c.alignment = WRAP
    style_header(ws, [11, 28, 62, 48, 62, 9])
    return ws


wb = Workbook()
ws = wb.active
ws.title = "Summary"
all_rows = api + ui
passed = sum(r["status"] == "PASS" for r in all_rows)
ws.append(["Sri Jayam ERP - Test Report"])
ws["A1"].font = Font(bold=True, size=16)
ws.append([f"Test cycle: {date.today():%d %B %Y}"])
ws.append([])
facts = [
    ("Result", f"{passed} of {len(all_rows)} tests passed" + ("" if passed == len(all_rows) else f" ({len(all_rows) - passed} failed - see the test sheets)")),
    ("Application", "Sri Jayam ERP billing system (web app: order entry, quotations, invoices & payments, job cards, delivery challans, customers, purchases & vendors, inventory, cash ledger & cheques, reports, pricing setup, users & roles)"),
    ("How it was tested", "Automated tests drive the real application the way staff do: every API action directly, and every screen in a real Chrome browser (typing, clicking, downloading). Each order is also priced independently from the raw rate tables and compared with what the system shows."),
    ("Test data", "14 customers, 33 orders (every one of the 14 services), 4 quotations, invoices that are paid / part-paid / unpaid (GST and cash bills), job cards in every stage, delivery challans, vendors, stock (3 items below reorder level), purchases, cash ledger entries and cheques, plus staff logins for each role."),
    ("Environment", "Identical application code and the same demo data as the UAT site. Tests that create and delete records ran on a separate copy, so the UAT site's data was left untouched."),
    ("Not covered", "Reading a supplier bill with the real AI service (needs the client's API key; the screen was tested for the case where no key is set), printing on physical printers, performance with many simultaneous users, WhatsApp / e-mail sending (not part of this release), and the correctness of the estimated rates against the client's own rate card."),
    ("Defects", f"{len(DEFECTS)} issues were found during this cycle and all were fixed and re-tested (see 'Issues found & fixed')."),
]
for k, v in facts:
    ws.append([k, v])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.cell(ws.max_row, 2).alignment = WRAP
ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 14
ws.column_dimensions["C"].width = 14
ws.column_dimensions["D"].width = 14
ws.merge_cells  # (layout kept simple: long text lives in column B)
ws.column_dimensions["B"].width = 110
ws.append([])
ws.append(["Area", "Tests", "Passed"])
for c in ws[ws.max_row]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = HEAD
by = OrderedDict()
for r in all_rows:
    by.setdefault(r["module"], [0, 0])
    by[r["module"]][0] += 1
    by[r["module"]][1] += r["status"] == "PASS"
for m, (n, p) in by.items():
    ws.append([m, n, p])
    ws.cell(ws.max_row, 3).fill = GREEN if n == p else RED
ws.append(["TOTAL", len(all_rows), passed])
for c in ws[ws.max_row]:
    c.font = Font(bold=True)
    c.fill = GREY

ws = wb.create_sheet("Issues found & fixed")
ws.append(["#", "Severity", "Where", "What was wrong", "What was done", "Status"])
for i, (sev, where, what, fix) in enumerate(DEFECTS, 1):
    ws.append([i, sev, where, what, fix, "Fixed & re-tested"])
    ws.cell(ws.max_row, 6).fill = GREEN
    for c in ws[ws.max_row]:
        c.alignment = WRAP
style_header(ws, [5, 10, 34, 70, 70, 18])

result_sheet(wb, "Browser tests (screens)", ui)
result_sheet(wb, "System tests (API)", api)

if shots and os.path.isdir(shots):
    ws = wb.create_sheet("Screenshots")
    ws.append(["Screens captured during the browser tests"])
    ws["A1"].font = Font(bold=True, size=13)
    row = 3
    KEEP = ["login-page", "dashboard-start-order", "order-form-visiting-card", "service-xerox-photocopy", "service-bill-books-stationery", "invoice-detail",
            "invoice-paid", "dashboard-live-board", "job-card-sheet", "delivery-challan-saved", "quotations-list", "pricing-setup", "price-matrix",
            "purchase-added", "cash-ledger", "reports", "layout-1366x768", "layout-menu-390x844",
            "company-details", "hsn-codes", "printing-and-tax", "order-gst-breakdown", "bill-book-wizard", "bill-book-copies", "outsource-list",
            "purchase-upload", "financial-statement", "quotation-estimated-delivery"]
    for k in KEEP:
        f = k + ".png"
        if not os.path.exists(os.path.join(shots, f)):
            continue
        img = XlImage(os.path.join(shots, f))
        scale = 520 / img.width
        img.width, img.height = 520, int(img.height * scale)
        ws.cell(row, 1, f[:-4].replace("-", " ").title()).font = Font(bold=True)
        ws.add_image(img, f"A{row + 1}")
        row += int(img.height / 20) + 4

wb.save(out)
print(f"saved {out}: {passed}/{len(all_rows)} passed")
